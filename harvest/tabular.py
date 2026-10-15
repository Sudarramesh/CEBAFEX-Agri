"""Bounded CSV/XLSX ingestion and reviewable heuristic column interpretation."""
from io import BytesIO
from difflib import SequenceMatcher
import re,zipfile
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
XML_ERRORS=(ParseError,)
try:
    from lxml.etree import XMLSyntaxError
    XML_ERRORS+=(XMLSyntaxError,)
except ImportError:
    pass
from .ingestion import read_csv_bytes,MAX_BYTES,MAX_RECORDS
from .data import ALIASES,BUYER_ALIASES,suggest_mapping

EXTRA={
 'crop':['commodity','produce','item name','vegetable name','crop description'],
 'available_kg':['stock kg','stock quantity','quantity in stock','current stock','quantity available kg','inventory kg','qoh'],
 'estimated_demand_kg':['estimated need kg','expected demand kg','demand quantity','recorded demand','customer demand kg'],
 'market_price_per_kg':['market rate','market price kg','unit market price','selling rate'],
 'harvest_age_days':['days since harvest','age days','produce age'],
 'buyer_name':['customer name','purchaser name','organisation name'],
 'buyer_type':['customer category','purchaser type','business type'],
 'required_kg':['requested kg','requirement kg','purchase quantity','requested quantity'],
 'offer_price_per_kg':['quoted price kg','buyer offer','purchase price','offered rate'],
 'transport_cost':['freight cost','delivery charge','transportation quote','shipping cost'],
}

def normalize(value):return re.sub(r'[^\w]+',' ',str(value).casefold()).strip()

def interpret_columns(frame,kind):
    original=suggest_mapping(frame.columns,kind);base=ALIASES if kind=='agriculture' else BUYER_ALIASES
    results=[];mapping=dict(original)
    for target in base:
        exact=[c for c in frame.columns if str(c).strip().casefold() in base[target]]
        source=original[target];method='Known alias';confidence='High'
        if source is None:
            if len(exact)>1:method='Ambiguous aliases';confidence='Review required'
            else:
                vocabulary=[normalize(x) for x in (*base[target],*EXTRA.get(target,[]))]
                scored=[]
                for column in frame.columns:
                    name=normalize(column)
                    # Forecast columns must never silently become entered demand estimates.
                    if target=='estimated_demand_kg' and 'forecast' in name:continue
                    score=max(SequenceMatcher(None,name,word).ratio() for word in vocabulary)
                    scored.append((score,str(column),column))
                scored.sort(reverse=True,key=lambda v:(v[0],v[1]))
                if scored and scored[0][0]>=.74 and (len(scored)==1 or scored[0][0]-scored[1][0]>=.12):
                    source=scored[0][2];method='Heuristic suggestion';confidence='Review required'
                else:method='Manual interpretation needed';confidence='Review required'
        mapping[target]=source
        results.append({'target':target,'source':source,'method':method,'confidence':confidence})
    # Avoid auto-assigning a single source to multiple meanings.
    for source in [s for s in mapping.values() if s is not None]:
        targets=[t for t,v in mapping.items() if v==source]
        if len(targets)>1:
            for row in results:
                if row['target'] in targets:
                    row.update(source=None,method='Ambiguous aliases',confidence='Review required');mapping[row['target']]=None
    return mapping,results


def _workbook(payload):
    if not payload or len(payload)>MAX_BYTES:raise ValueError('File exceeds 10 MB or is empty')
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(x.file_size for x in archive.infolist())>50*1024*1024:raise ValueError('Expanded workbook exceeds 50 MB')
            if len(archive.infolist())>2000:raise ValueError('Workbook has too many internal entries')
        return load_workbook(BytesIO(payload),read_only=True,data_only=False,keep_links=False)
    except (zipfile.BadZipFile,KeyError,OSError,InvalidFileException,EOFError,*XML_ERRORS) as exc:raise ValueError('Invalid Excel workbook') from exc


def sheets(payload):
    book=_workbook(payload)
    try:return list(book.sheetnames)
    finally:book.close()


def read_table(payload,filename,sheet=None):
    if str(filename).lower().endswith('.csv'):return read_csv_bytes(payload)
    if not str(filename).lower().endswith('.xlsx'):raise ValueError('Use a CSV or XLSX file')
    book=_workbook(payload)
    try:
        selected=book[sheet or book.sheetnames[0]]
        if selected.max_row and selected.max_row>MAX_RECORDS+1:raise ValueError('Workbook exceeds 50000 rows')
        if selected.max_column and selected.max_column>200:raise ValueError('Workbook exceeds 200 columns')
        rows=[]
        for record in selected.iter_rows():
            values=[]
            for cell in record:
                if cell.data_type=='f':raise ValueError('Excel formulas are not accepted; export values first')
                value=cell.value
                if hasattr(value,'isoformat'):value=value.isoformat()
                values.append(value)
            rows.append(values)
            if len(rows)>MAX_RECORDS+1:raise ValueError('Workbook exceeds 50000 rows')
        if not rows:raise ValueError('Workbook contains no rows')
        headers=[str(v).strip() if v is not None else '' for v in rows[0]]
        if any(not h for h in headers) or len({h.casefold() for h in headers})!=len(headers):raise ValueError('Column headers must be nonblank and unique')
        return pd.DataFrame(rows[1:],columns=headers)
    except (zipfile.BadZipFile,KeyError,OSError,InvalidFileException,EOFError,*XML_ERRORS) as exc:
        raise ValueError('Invalid Excel workbook') from exc
    finally:book.close()
